# -*- coding: utf-8 -*-
"""
일일결산 양식을 만든다.

  python scripts/일일결산_양식만들기.py 2026 10
  (연·월을 안 주면 06_매출 폴더의 마지막 달 다음 달로 만든다)

■ 한 줄 = 한 손님
  손님이 테와 누진렌즈를 같이 사도 한 줄에 적는다.
  품목 칸이 안경테 / 단초점 / 누진 / 콘택트 네 개라서,
  산 것만 드롭다운으로 고르면 된다.

■ 고르는 것이 곧 건수다
  누진 칸에 제품을 고르면 그날 누진 1건으로 자동으로 세어진다.
  따로 숫자를 적을 필요가 없다.

■ 금액은 품목별로 나누지 않는다
  한 손님이 낸 돈을 테 얼마 / 렌즈 얼마로 쪼갤 수 없기 때문이다.
  대신 '그 품목이 들어간 판매의 합계와 평균' 을 월합계에서 보여준다.

■ 제품 목록은 설정 시트에서 늘린다
  설정 시트 아래에 한 줄 추가하면 드롭다운에 바로 나온다. (수식 고칠 필요 없음)
  처음 목록은 03_데이터/사입단가.xlsx 에서 가져온다.

■ 구조
  월합계 — 일자별 표 + 누진·콘택트 건수 요약 (전부 수식)
  1일~31일 — 하루치 판매. 맨 위에 오늘 판매액과 누진·콘택트 건수가 뜬다.
  설정   — 담당자와 제품 목록
"""
import calendar
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).resolve().parent.parent
SALES_DIR = BASE / "06_매출"
PRICE = BASE / "03_데이터" / "사입단가.xlsx"

STAFF = ["이철기", "김상기", "한나영"]
# 품목 칸. 이 칸에 제품을 고르면 그게 곧 1건이다.
ITEMS = ["안경테", "단초점", "누진", "콘택트"]
# 사입단가표의 '구분' 중 어느 것을 각 칸의 목록으로 쓸지
SEED = {"안경테": "프레임", "단초점": "렌즈", "누진": None, "콘택트": "콘택트"}
KEY = ["누진", "콘택트"]        # 특별히 건수를 따로 보고 싶은 품목
MONEY = ["현금", "현금영수증", "상품권", "카드", "미수"]
ROWS_PER_DAY = 40
HDR = 5
FIRST = HDR + 1
LIST_ROWS = 400                # 제품을 이만큼까지 늘릴 수 있다
WEEK = ["월", "화", "수", "목", "금", "토", "일"]

COLS = ([("번호", 5), ("담당자", 10), ("고객명", 12)]
        + [(i, 15 if i in ("안경테", "콘택트") else 12) for i in ITEMS]
        + [("상세", 16)]
        + [(m, 12 if m == "현금영수증" else 11) for m in MONEY]
        + [("합계", 13), ("비고", 14)])
A = {n: get_column_letter(i) for i, (n, _) in enumerate(COLS, 1)}

NAVY = PatternFill("solid", fgColor="1F3864")
ITEMFILL = PatternFill("solid", fgColor="4472C4")
KEYFILL = PatternFill("solid", fgColor="C55A11")     # 누진·콘택트는 눈에 띄게
PAY = PatternFill("solid", fgColor="2E5C8A")
AUTO = PatternFill("solid", fgColor="EDEDED")
TOTAL = PatternFill("solid", fgColor="FFF2CC")
BIG = PatternFill("solid", fgColor="DDEBF7")
WHITE = Font(color="FFFFFF", bold=True, size=10)
BOX = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def head_fill(name):
    if name in KEY:
        return KEYFILL
    if name in ITEMS:
        return ITEMFILL
    if name in MONEY:
        return PAY
    return NAVY


def load_products():
    """사입단가.xlsx 에서 제품 이름을 가져온다. 없으면 빈 목록."""
    out = {k: [] for k in ITEMS}
    if not PRICE.exists():
        return out, False
    try:
        wb = openpyxl.load_workbook(PRICE, data_only=True)
    except Exception:
        return out, False
    ws = wb["사입단가"] if "사입단가" in wb.sheetnames else wb.active
    by_kind = {}
    for kind, name in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        if kind and name:
            by_kind.setdefault(str(kind).strip(), []).append(str(name).strip())
    wb.close()
    for col, kind in SEED.items():
        if kind and kind in by_kind:
            out[col] = sorted(set(by_kind[kind]))   # 같은 이름 중복 제거
    return out, True


def build_day(wb, y, m, day):
    ws = wb.create_sheet(f"{day}일")
    last, t = FIRST + ROWS_PER_DAY - 1, FIRST + ROWS_PER_DAY

    ws["A1"] = f"{m}월 {day}일 ({WEEK[calendar.weekday(y, m, day)]})"
    ws["A1"].font = Font(bold=True, size=16)
    ws[f'{A["안경테"]}1'] = ("한 줄에 손님 한 분.  산 것만 골라주세요 "
                             "(테+누진이면 두 칸).  돈은 받은 방법에 한 번만.")
    ws[f'{A["안경테"]}1'].font = Font(size=10, color="808080")

    ws["A3"] = "오늘 판매"
    ws["A3"].font = Font(bold=True, size=11, color="1F3864")
    # 18pt 숫자는 한 칸에 안 들어가 ##### 로 잘린다. 두 칸을 합쳐 쓴다.
    ws.merge_cells("B3:C3")
    ws["B3"] = f'={A["합계"]}{t}'
    ws["B3"].font = Font(bold=True, size=18, color="1F3864")
    ws["B3"].number_format = '#,##0"원"'
    ws["B3"].alignment = Alignment(horizontal="left", vertical="center")
    for it in ITEMS:
        c = ws[f"{A[it]}3"]
        c.value = f'=IF({A[it]}{t}=0,"",{A[it]}{t}&"건")'
        c.font = Font(bold=True, size=13 if it in KEY else 11,
                      color="C55A11" if it in KEY else "1F3864")
        c.alignment = Alignment(horizontal="center")
    ws[f'{A["미수"]}3'] = (f'=IF({A["미수"]}{t}=0,"",'
                           f'"미수 "&TEXT({A["미수"]}{t},"#,##0"))')
    ws[f'{A["미수"]}3'].font = Font(bold=True, size=10, color="C00000")
    for i in range(1, len(COLS) + 1):
        ws.cell(3, i).fill = BIG

    for i, (name, _) in enumerate(COLS, 1):
        c = ws.cell(HDR, i, name)
        c.fill = head_fill(name)
        c.font = WHITE
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX

    for n, r in enumerate(range(FIRST, last + 1), 1):
        ws.cell(r, 1, n).alignment = Alignment(horizontal="center")
        c = ws.cell(r, list(A).index("합계") + 1,
                    f'=IF(SUM({A["현금"]}{r}:{A["카드"]}{r})=0,"",'
                    f'SUM({A["현금"]}{r}:{A["카드"]}{r}))')
        c.fill = AUTO
        for i in range(1, len(COLS) + 1):
            ws.cell(r, i).border = BOX

    # 합계 줄 — 품목은 '고른 칸의 개수', 돈은 합계
    ws.cell(t, 1, "합계").font = Font(bold=True)
    for it in ITEMS:
        col = A[it]
        ws.cell(t, list(A).index(it) + 1,
                f'=COUNTIF({col}{FIRST}:{col}{last},"?*")').font = Font(bold=True)
    for name in MONEY + ["합계"]:
        col = A[name]
        ws.cell(t, list(A).index(name) + 1,
                f"=SUM({col}{FIRST}:{col}{last})").font = Font(bold=True)
    for i in range(1, len(COLS) + 1):
        ws.cell(t, i).fill = TOTAL
        ws.cell(t, i).border = BOX

    for i, (name, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for name in MONEY + ["합계"]:
        col = list(A).index(name) + 1
        for r in range(FIRST, t + 1):
            ws.cell(r, col).number_format = "#,##0"
    ws.freeze_panes = f"D{FIRST}"

    # showDropDown 은 이름과 반대로 '화살표를 숨길지' 다. False 여야 목록이 보인다.
    warn = dict(errorStyle="warning", errorTitle="목록에 없는 이름입니다",
                error="목록에 없어도 그대로 쓰시려면 '예' 를 누르세요.\n"
                      "자주 쓰실 이름이면 설정 시트에 추가해두세요.",
                showErrorMessage=True, showDropDown=False)
    dv = DataValidation(type="list", formula1="담당자목록", allow_blank=True, **warn)
    ws.add_data_validation(dv)
    dv.add(f'{A["담당자"]}{FIRST}:{A["담당자"]}{last}')
    for it in ITEMS:
        d = DataValidation(type="list", formula1=f"{it}목록", allow_blank=True, **warn)
        ws.add_data_validation(d)
        d.add(f"{A[it]}{FIRST}:{A[it]}{last}")


def build_setup(wb, products, found):
    st = wb.create_sheet("설정")
    st["A1"] = "여기에 이름을 추가하면 일별 시트의 드롭다운에 바로 나옵니다."
    st["A1"].font = Font(bold=True, size=12, color="C00000")
    st["A2"] = ("각 칸 아래로 계속 이어서 적으시면 됩니다. 수식은 안 고치셔도 됩니다. "
                f"(칸마다 {LIST_ROWS}개까지)")
    st["A2"].font = Font(size=10, color="808080")

    heads = [("담당자", 1)] + [(it, 3 + n) for n, it in enumerate(ITEMS)]
    for name, ci in heads:
        c = st.cell(3, ci, name)
        c.fill = KEYFILL if name in KEY else NAVY
        c.font = WHITE
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
    for n, s in enumerate(STAFF):
        st.cell(4 + n, 1, s).border = BOX
    for name, ci in heads[1:]:
        for n, p in enumerate(products[name]):
            st.cell(4 + n, ci, p).border = BOX

    # 빈 칸 안내는 목록 밖(오른쪽)에 적는다. 목록 안에 두면 드롭다운에 딸려 나온다.
    empty = [n for n, _ in heads[1:] if not products[n]]
    side = 4 + len(ITEMS)
    if empty:
        c = st.cell(4, side, f"← {' · '.join(empty)} 칸이 비어 있습니다. "
                             f"취급하시는 제품 이름을 적어주세요.")
        c.font = Font(size=10, bold=True, color="C00000")

    st.column_dimensions["A"].width = 14
    for name, ci in heads[1:]:
        st.column_dimensions[get_column_letter(ci)].width = 17
    st.column_dimensions[get_column_letter(side)].width = 4

    note = 4 + LIST_ROWS + 2
    for line in [
        "제품 목록 — 쓰는 법",
        "",
        "1. 위 표의 각 칸 맨 아래에 한 줄 추가하시면 드롭다운에 바로 나옵니다.",
        "2. 누진 칸은 처음엔 비어 있습니다. 취급하시는 누진 제품 이름을 적어주세요.",
        "   (사입단가표의 렌즈 목록에 누진 제품이 따로 없어서 비워뒀습니다)",
        "3. 목록에 없는 이름을 그냥 타이핑하셔도 됩니다. 확인 창에서 '예' 를 누르면 됩니다.",
        "4. 이름은 세 명이 똑같이 써야 집계가 맞습니다. 되도록 드롭다운에서 고르세요.",
        "",
        "※ 누진 칸에 제품을 고르시면 그게 곧 '누진 1건' 으로 자동으로 세어집니다.",
        "   따로 개수를 적으실 필요 없습니다.",
    ]:
        st.cell(note, 1, line)
        note += 1
    st.cell(4 + LIST_ROWS + 2, 1).font = Font(bold=True, size=12)

    # 목록이 늘어나도 따라오도록 OFFSET 으로 잡는다. 비어 있어도 오류가 안 나게 MAX(1,..)
    # ※ attr_text 에 '=' 를 붙이면 안 된다. 엑셀이 이름을 통째로 못 읽어
    #   드롭다운이 조용히 사라진다. (2026-09-06 에 이걸로 한 번 깨졌다)
    def dyn(col):
        return (f"OFFSET(설정!${col}$4,0,0,"
                f"MAX(1,COUNTA(설정!${col}$4:${col}${3+LIST_ROWS})),1)")
    wb.defined_names.add(DefinedName("담당자목록", attr_text=dyn("A")))
    for name, ci in heads[1:]:
        wb.defined_names.add(
            DefinedName(f"{name}목록", attr_text=dyn(get_column_letter(ci))))


def build_summary(wb, y, m, days):
    t = FIRST + ROWS_PER_DAY
    su = wb.create_sheet("월합계")
    su["A1"] = f"{y}년 {m}월"
    su["A1"].font = Font(bold=True, size=18)

    # 맨 위에 누진·콘택트 건수를 크게. 칸을 합쳐 써야 ##### 로 안 잘린다.
    su["D1"] = "이번 달"
    su["D1"].font = Font(bold=True, size=11, color="808080")
    su.merge_cells("E1:J1")
    su["E1"].font = Font(bold=True, size=20, color="C55A11")
    su["E1"].alignment = Alignment(horizontal="left", vertical="center")

    head = ["일자", "손님수"] + ITEMS + MONEY + ["판매합계"]
    for i, h in enumerate(head, 1):
        c = su.cell(3, i, h)
        c.fill = head_fill(h) if h != "판매합계" else NAVY
        c.font = WHITE
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
    for d in range(1, days + 1):
        r = 3 + d
        su.cell(r, 1, f"{d}일")
        su.cell(r, 2, f"=COUNTIF('{d}일'!{A['합계']}${FIRST}:"
                      f"{A['합계']}${t-1},\">0\")")
        for i, name in enumerate(ITEMS + MONEY + ["합계"], 3):
            su.cell(r, i, f"='{d}일'!{A[name]}{t}")
        for i in range(1, len(head) + 1):
            su.cell(r, i).border = BOX
    tr = 4 + days
    su.cell(tr, 1, "합계").font = Font(bold=True)
    for i in range(2, len(head) + 1):
        col = get_column_letter(i)
        su.cell(tr, i, f"=SUM({col}4:{col}{3+days})").font = Font(bold=True)
    for i in range(1, len(head) + 1):
        su.cell(tr, i).fill = TOTAL
        su.cell(tr, i).border = BOX
    su["E1"] = "=" + '&"   ·   "&'.join(
        f'"{k} "&{get_column_letter(3 + ITEMS.index(k))}{tr}&"건"' for k in KEY)

    # 품목별 요약
    ir = tr + 2
    su.cell(ir, 1, "품목별").font = Font(bold=True, size=13)
    su.cell(ir, 3, "※ 한 손님이 테와 누진을 같이 사면 각각 1건으로 셉니다.")
    su.cell(ir, 3).font = Font(size=9, color="808080")
    su.cell(ir + 1, 3, "   금액은 그 품목이 들어간 판매의 합계라, 다 더하면 총매출보다 큽니다.")
    su.cell(ir + 1, 3).font = Font(size=9, color="808080")
    for i, h in enumerate(["품목", "건수", "이 품목이 든 판매액", "평균 객단가"], 1):
        c = su.cell(ir + 2, i, h)
        c.fill = NAVY
        c.font = WHITE
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
    for n, it in enumerate(ITEMS):
        r = ir + 3 + n
        c = su.cell(r, 1, it)
        if it in KEY:
            c.font = Font(bold=True, color="C55A11")
        icol, kcol = A[it], A["합계"]
        su.cell(r, 2, "=" + "+".join(f"'{d}일'!{icol}{t}" for d in range(1, days + 1)))
        su.cell(r, 3, "=" + "+".join(
            f'SUMIF(\'{d}일\'!{icol}${FIRST}:{icol}${t-1},"?*",'
            f"'{d}일'!{kcol}${FIRST}:{kcol}${t-1})" for d in range(1, days + 1)))
        # IFERROR 를 쓰면 엑셀이 _xlfn.IFERROR 라는 깨진 이름을 하나 만든다.
        su.cell(r, 4, f'=IF(B{r}>0,C{r}/B{r},"")')
        for i in range(1, 5):
            su.cell(r, i).border = BOX

    widths = [10, 8] + [11] * len(ITEMS) + [11] * len(MONEY) + [14]
    for i, w in enumerate(widths, 1):
        su.column_dimensions[get_column_letter(i)].width = w
    su.column_dimensions["A"].width = 14
    su.column_dimensions["C"].width = 20
    su.column_dimensions["D"].width = 14
    money_start = 3 + len(ITEMS)
    for row in su.iter_rows(min_row=4, min_col=money_start, max_col=len(head)):
        for c in row:
            c.number_format = "#,##0"
    for row in su.iter_rows(min_row=ir + 3, min_col=3, max_col=4):
        for c in row:
            c.number_format = "#,##0"
    su.freeze_panes = "C4"


def make_workbook(y, m, extra=None):
    """빈 양식 워크북을 만들어 돌려준다. (일지_옮기기.py 도 이걸 쓴다)

    extra: {"누진": ["대명누진", ...]} 처럼 목록에 더 넣을 이름.
    """
    products, found = load_products()
    for k, names in (extra or {}).items():
        products[k] = sorted(set(products.get(k, [])) | set(n for n in names if n))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    days = calendar.monthrange(y, m)[1]
    build_summary(wb, y, m, days)
    for d in range(1, days + 1):
        build_day(wb, y, m, d)
    build_setup(wb, products, found)
    return wb, products, found


def next_month():
    months = []
    for f in SALES_DIR.glob("*.xlsx"):
        mm = re.search(r"(\d{2,4})\s*[년\-.]\s*(\d{1,2})\s*월?", f.stem)
        if mm:
            yy = int(mm.group(1))
            months.append((yy + 2000 if yy < 100 else yy, int(mm.group(2))))
    y, m = max(months) if months else (2026, 10)
    m += 1
    return (y + 1, 1) if m == 13 else (y, m)


def main():
    if len(sys.argv) >= 3:
        y, m = int(sys.argv[1]), int(sys.argv[2])
    else:
        y, m = next_month()

    wb, products, found = make_workbook(y, m)
    days = calendar.monthrange(y, m)[1]

    out = SALES_DIR / f"{y}년 {m}월.xlsx"
    if out.exists():
        out = SALES_DIR / f"{y}년 {m}월 (새양식).xlsx"
    wb.save(out)

    print(f"만들었습니다: {out.name}")
    print(f"   시트 — 월합계 · 1일~{days}일 · 설정")
    print(f"   품목 칸 — {' / '.join(ITEMS)}  (드롭다운에서 고르면 그게 곧 1건)")
    if found:
        for it in ITEMS:
            n = len(products[it])
            print(f"      {it:<5} 제품 {n:>3}개" +
                  ("   ← 설정 시트에 적어주세요" if n == 0 else ""))
    else:
        print("      [!] 사입단가.xlsx 를 못 읽어 제품 목록이 비어 있습니다")


if __name__ == "__main__":
    main()
